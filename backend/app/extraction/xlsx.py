"""XLSX text extraction via openpyxl: every non-empty cell, sheet by sheet.

An empty workbook is not an OCR case — it yields an empty ExtractedDocument
(char_count 0) and the pipeline decides what to do with it.
"""

import io

from app.extraction.base import ExtractedDocument, PageText, ParserUnavailable
from app.extraction.sniff import MIME_XLSX


class XlsxHandler:
    """Extracts cell values as lines; None cells are skipped entirely."""

    def extract(self, data: bytes) -> ExtractedDocument:
        try:
            import openpyxl
        except ImportError as exc:
            raise ParserUnavailable("openpyxl is not installed") from exc
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    lines.extend(str(cell) for cell in row if cell is not None)
        finally:
            workbook.close()  # read_only mode holds a file handle until closed
        text = "\n".join(lines)
        return ExtractedDocument(
            text=text,
            pages=[PageText(page_no=1, text=text)],
            mime_sniffed=MIME_XLSX,
            char_count=len(text),
            ocr_used=False,
        )
