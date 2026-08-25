"""Synthetic corpus generator for the DMS classifiers.

Generates N fully synthetic records spread over doc_type x level combos, honours the
per-level entity-count SPECS, and renders every record as .docx, .xlsx and .pdf.
Emits manifest.csv (one row per record) plus a manifest.json sidecar carrying the
per-record source text used at generation time (the exporter needs it because PDF
text is not read back without an extra dependency).

Deterministic for a given --seed: fixed RNG/Faker consumption order, uuid5 ids,
fixed generated_at base. All data is synthetic (Faker 'en_PK'); nothing real.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
import uuid
from collections import Counter
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from faker import Faker
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook

from entities import (
    CARD_RANGE_BY_LEVEL,
    SPECS,
    make_account,
    make_card,
    make_cnic,
    make_passport,
)
from templates import LABEL_PHRASES, TEMPLATES, RenderContext, strip_label_phrases

DOC_TYPES: tuple[str, ...] = tuple(TEMPLATES)
LEVELS: tuple[str, ...] = ("internal", "confidential", "restricted")
# Passports are outside spec §3.7; phase-1 choice: rare, restricted-only.
PASSPORT_RANGE_BY_LEVEL: dict[str, tuple[int, int]] = {
    "restricted": (0, 1),
    "confidential": (0, 0),
    "internal": (0, 0),
}
BASE_GENERATED_AT = datetime(2026, 1, 1, tzinfo=UTC)
NAMES_PER_RECORD = 8
MANIFEST_COLUMNS = ("file_path", "label_doc_type", "label_level", "version_id", "generated_at")


def _build_facts(rng: random.Random, fake: Faker, level: str) -> tuple[list[str], dict[str, int]]:
    """Build the entity/salary fact lines for one record exactly per SPECS."""
    spec = SPECS[level]
    names = [fake.name() for _ in range(NAMES_PER_RECORD)]
    # Draw every count exactly once so the sidecar counts always match the emitted lines.
    cnic_count = rng.randint(*spec["cnic"])
    card_count = rng.randint(*CARD_RANGE_BY_LEVEL[level])
    passport_count = rng.randint(*PASSPORT_RANGE_BY_LEVEL[level])
    account_count = rng.randint(*spec["account"])
    salary_count = rng.randint(*spec["salary"])
    facts = [
        *(f"CNIC on record: {make_cnic(rng)}" for _ in range(cnic_count)),
        *(f"Card token held for billing: {make_card(rng)}" for _ in range(card_count)),
        *(f"Passport reference: {make_passport(rng)}" for _ in range(passport_count)),
        *(f"Settlement account reference: {make_account(rng)}" for _ in range(account_count)),
        *(
            f"{name}: monthly remuneration PKR {rng.randint(25_000, 500_000):,}"
            for name in names[:salary_count]
        ),
    ]
    counts = {
        "cnic": cnic_count,
        "card": card_count,
        "passport": passport_count,
        "account": account_count,
        "salary": salary_count,
    }
    return facts, counts


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)


def _write_xlsx(path: Path, lines: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("line_no", "text"))
    for number, line in enumerate(lines, start=1):
        sheet.append((number, line))
    workbook.save(path)


def _write_pdf(path: Path, lines: list[str]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        safe = line.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(w=0, h=5, text=safe, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(path)


def generate_corpus(count: int, out_dir: Path, seed: int) -> None:
    """Generate ``count`` synthetic records (3 renderings each) plus manifests into *out_dir*."""
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    fake = Faker("en_PK")
    fake.seed_instance(seed)

    csv_rows: list[dict[str, str]] = []
    json_records: list[dict[str, Any]] = []
    leaks_rewritten = 0

    for index in range(count):
        combo = index % (len(DOC_TYPES) * len(LEVELS))
        doc_type = DOC_TYPES[combo // len(LEVELS)]
        level = LEVELS[combo % len(LEVELS)]

        facts, entity_counts = _build_facts(rng, fake, level)
        ctx = RenderContext(
            doc_type=doc_type,
            level=level,
            company=fake.company(),
            person_names=tuple(fake.name() for _ in range(2)),
            facts=tuple(facts),
        )
        sections = TEMPLATES[doc_type](rng, ctx)
        raw_body = "\n".join(sections)
        body = strip_label_phrases(raw_body)
        if body != raw_body:
            leaks_rewritten += sum(
                1
                for phrase in LABEL_PHRASES
                if re.search(re.escape(phrase), raw_body, re.IGNORECASE)
            )

        version_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"dms-ml/{seed}/{index}"))
        generated_at = (BASE_GENERATED_AT + timedelta(seconds=index)).isoformat()
        stems = {ext: f"record_{index:04d}.{ext}" for ext in ("docx", "xlsx", "pdf")}
        _write_docx(out_dir / stems["docx"], body.splitlines())
        _write_xlsx(out_dir / stems["xlsx"], body.splitlines())
        _write_pdf(out_dir / stems["pdf"], body.splitlines())

        csv_rows.append(
            {
                "file_path": stems["docx"],
                "label_doc_type": doc_type,
                "label_level": level.capitalize(),
                "version_id": version_id,
                "generated_at": generated_at,
            }
        )
        json_records.append(
            {
                "index": index,
                "version_id": version_id,
                "doc_type": doc_type,
                "level": level,
                "generated_at": generated_at,
                "files": stems,
                "source_text": body,
                "entity_counts": entity_counts,
            }
        )

    with (out_dir / "manifest.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(MANIFEST_COLUMNS))
        writer.writeheader()
        writer.writerows(csv_rows)

    sidecar = {"seed": seed, "count": count, "records": json_records}
    (out_dir / "manifest.json").write_text(
        json.dumps(sidecar, indent=2, sort_keys=True), encoding="utf-8"
    )

    distribution = Counter((r["label_doc_type"], r["label_level"]) for r in csv_rows)
    print(f"Generated {count} synthetic records ({count * 3} files) into {out_dir}")
    print(f"Label-phrase violations rewritten post-check: {leaks_rewritten}")
    print("Distribution (doc_type, level -> count):")
    for (doc_type, level), n in sorted(distribution.items()):
        print(f"  {doc_type}, {level} -> {n}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate the synthetic DMS training corpus.")
    parser.add_argument("--count", type=int, default=50, help="number of records (3 files each)")
    parser.add_argument("--out", type=Path, required=True, help="output directory")
    parser.add_argument("--seed", type=int, default=42, help="deterministic seed")
    args = parser.parse_args(argv)
    generate_corpus(count=args.count, out_dir=args.out, seed=args.seed)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
